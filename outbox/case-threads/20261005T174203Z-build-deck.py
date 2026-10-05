#!/usr/bin/env python3
"""Build the 2-slide NGM savings review deck for Fabian & Christoph.

Reads data from 20261005T174203Z-deck-data.json (same dir) so verified
numbers can be patched without touching the layout code.
Output: 20261005T174203Z-ngm-savings-review.pptx
"""
import json
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = json.load(open(os.path.join(HERE, "20261005T174203Z-deck-data.json")))

# --- palette (CAST AI flavored, Siemens-safe) ---
INK      = RGBColor(0x1B, 0x1B, 0x1F)
GREY     = RGBColor(0x5A, 0x5A, 0x66)
CAST     = RGBColor(0x00, 0xB3, 0x91)   # cast green-teal
CAST_DK  = RGBColor(0x00, 0x7A, 0x63)
AMBER    = RGBColor(0xE8, 0x8B, 0x00)
L_ROW    = RGBColor(0xF2, 0xF7, 0xF5)
L_HEAD   = RGBColor(0x00, 0x7A, 0x63)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LINE     = RGBColor(0xD5, 0xDE, 0xDB)
P1C      = RGBColor(0xC7, 0x2C, 0x41)
P2C      = RGBColor(0xE8, 0x8B, 0x00)
P3C      = RGBColor(0x4A, 0x6F, 0xCF)

prs = Presentation()
prs.slide_width  = Emu(12192000)   # 16:9
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_footer(slide, n):
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(7.02), Inches(12.3), Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = "Unrestricted | © Siemens 2026 | SI GSW CLO – NGM cost review | CAST AI"
    r.font.size = Pt(9); r.font.color.rgb = GREY
    r2 = p.add_run(); r2.text = f"    {n}"
    r2.font.size = Pt(9); r2.font.color.rgb = GREY


def add_title(slide, text, sub):
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(1.02))  # rectangle
    bar.fill.solid(); bar.fill.fore_color.rgb = CAST_DK; bar.line.fill.background()
    tf = bar.text_frame; tf.margin_left = Inches(0.55); tf.margin_top = Inches(0.10)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    r = p2.add_run(); r.text = sub
    r.font.size = Pt(12); r.font.color.rgb = RGBColor(0xD8, 0xF3, 0xEC)


def style_cell(cell, size=10, bold=False, color=INK, align=PP_ALIGN.LEFT, fill=None):
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


# ============================================================ SLIDE 1
s = prs.slides.add_slide(BLANK)
add_title(s, DATA["s1_title"], DATA["s1_subtitle"])

# headline chips
chips = DATA["s1_headline"]
cx = Inches(0.55)
for label, value, col in chips:
    colrgb = RGBColor.from_string(col.lstrip("#")) if isinstance(col, str) else col
    box = s.shapes.add_shape(1, cx, Inches(1.22), Inches(3.0), Inches(0.78))
    box.fill.solid(); box.fill.fore_color.rgb = L_ROW; box.line.color.rgb = LINE
    tf = box.text_frame; tf.margin_left = Inches(0.12); tf.margin_top = Inches(0.05)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = value
    r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = colrgb
    p2 = tf.add_paragraph(); r = p2.add_run(); r.text = label
    r.font.size = Pt(10); r.font.color.rgb = GREY
    cx += Inches(3.25)

# savings table (left)
rows = DATA["s1_rows"]
tbl_shape = s.shapes.add_table(len(rows) + 1, 3, Inches(0.55), Inches(2.25), Inches(5.9), Inches(0.4 * (len(rows) + 1)))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.5)
tbl.columns[1].width = Inches(1.7)
tbl.columns[2].width = Inches(1.7)
hdr = ["Cluster", "Gross savings", "Net savings*"]
for j, h in enumerate(hdr):
    c = tbl.cell(0, j); c.text = h
    style_cell(c, 11, True, WHITE, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, L_HEAD)
for i, row in enumerate(rows, start=1):
    fill = L_ROW if i % 2 == 0 else WHITE
    bold = row.get("total", False)
    for j, key in enumerate(["name", "gross", "net"]):
        c = tbl.cell(i, j); c.text = row[key]
        style_cell(c, 11, bold, INK, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, fill)

fn = s.shapes.add_textbox(Inches(0.55), Inches(2.25) + Inches(0.4 * (len(rows) + 1)) + Inches(0.06), Inches(5.9), Inches(0.8))
tf = fn.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = DATA["s1_footnote"]
r.font.size = Pt(9); r.font.color.rgb = GREY

# "why they don't see it" box (right)
box = s.shapes.add_shape(1, Inches(6.85), Inches(2.25), Inches(6.0), Inches(4.35))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFB, 0xF4, 0xE8); box.line.color.rgb = AMBER
tf = box.text_frame; tf.word_wrap = True
tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.15); tf.margin_top = Inches(0.12)
p = tf.paragraphs[0]; r = p.add_run(); r.text = DATA["s1_why_title"]
r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = RGBColor(0x9A, 0x60, 0x00)
for bullet in DATA["s1_why_bullets"]:
    bp = tf.add_paragraph(); bp.space_before = Pt(6)
    r = bp.add_run(); r.text = "▸ " + bullet[0]
    r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = INK
    sp = tf.add_paragraph()
    r = sp.add_run(); r.text = "  " + bullet[1]
    r.font.size = Pt(10); r.font.color.rgb = GREY

# speaker notes
notes = s.notes_slide.notes_text_frame
notes.text = DATA["s1_notes"]
add_footer(s, 1)

# ============================================================ SLIDE 2
s = prs.slides.add_slide(BLANK)
add_title(s, DATA["s2_title"], DATA["s2_subtitle"])

rows = DATA["s2_rows"]
ncols = 7
tbl_shape = s.shapes.add_table(len(rows) + 1, ncols, Inches(0.35), Inches(1.30), Inches(12.6), Inches(0.47 * (len(rows) + 1)))
tbl = tbl_shape.table
widths = [0.45, 2.30, 0.90, 3.20, 2.30, 1.75, 1.70]  # sums to exactly 12.6
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
hdr = ["Prio", "Recommendation", "Scope", "Expected impact", "Controls / guardrails", "Note", "$/mo (H · I · K → Σ)"]
for j, h in enumerate(hdr):
    c = tbl.cell(0, j); c.text = h
    style_cell(c, 10, True, WHITE, PP_ALIGN.LEFT, L_HEAD)

prio_col = {"P1": P1C, "P2": P2C, "P3": P3C}
for i, row in enumerate(rows, start=1):
    fill = L_ROW if i % 2 == 0 else WHITE
    if row.get("total"):
        fill = RGBColor(0xE4, 0xF1, 0xEE)
    vals = [row["prio"], row["rec"], row["scope"], row["impact"], row["controls"], row["note"], row["money"]]
    for j, v in enumerate(vals):
        c = tbl.cell(i, j); c.text = v
        if j == 0:
            style_cell(c, 10, True, prio_col.get(v, INK), PP_ALIGN.CENTER, fill)
        elif j == 6:
            style_cell(c, 9, row.get("total", False), CAST_DK if row.get("total") else INK, PP_ALIGN.LEFT, fill)
        else:
            style_cell(c, 9, row.get("total", False), INK, PP_ALIGN.LEFT, fill)

fn = s.shapes.add_textbox(Inches(0.35), Inches(6.62), Inches(12.6), Inches(0.38))
tf = fn.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = DATA["s2_footnote"]
r.font.size = Pt(9); r.font.color.rgb = GREY

notes = s.notes_slide.notes_text_frame
notes.text = DATA["s2_notes"]
add_footer(s, 2)

out = os.path.join(HERE, "20261005T174203Z-ngm-savings-review.pptx")
prs.save(out)
print("WROTE", out)
