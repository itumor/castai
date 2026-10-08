#!/usr/bin/env python3
"""Build customer mail DOCX: ORG CPS negative-savings reply, claim -> snapshot -> claim -> snapshot -> calculation."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SHOTS = '/Users/eramadan/castai/outbox/case-threads/console-shots-20261007T2155Z'
OUT = '/Users/eramadan/castai/outbox/case-threads/20261008T0122Z-cps-negative-savings-mail.docx'

doc = Document()
for section in doc.sections:
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)  # A4
    section.left_margin = section.right_margin = Inches(0.8)

style = doc.styles['Normal']
style.font.name = 'Calibri'; style.font.size = Pt(10.5)

def p(text='', bold=False, italic=False, size=10.5, space_after=6, align=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold; run.italic = italic; run.font.size = Pt(size)
    par.paragraph_format.space_after = Pt(space_after)
    if align: par.alignment = align
    return par

def heading(text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = True; run.font.size = Pt(13)
    par.paragraph_format.space_before = Pt(12); par.paragraph_format.space_after = Pt(6)

def figure(path, width, caption):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    run = cap.add_run(caption)
    run.italic = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)

# ---------- header ----------
p('Subject: RE: ORG CPS — two clusters with negative savings', bold=True, size=11)
p('To: Sergej Petrovski <sergej.petrovski@siemens.com>', size=9, space_after=0)
p('Cc: Lars Marin (tecracer) | From: Ebrahim Ramadan <ebrahim@cast.ai> | Date: 2026-10-08', size=9)
p('—' * 46, size=9)

p('Hello Sergej,')
p('short answer first: nothing is wrong — neither cluster is in minus. I verified this today, claim by claim, '
  'against your own console and our API. Claims and proofs below.')

# ---------- claim 1 ----------
heading('Claim 1 — Both clusters have positive realized savings. Right now, every day.')
p('Proof: the Savings Report of ORG CPS as of this morning (October, org scope). '
  'Baseline spend $6,841 minus actual spend $3,849 = realized savings +$2,992 in the first week of October alone. '
  'A cluster losing money cannot produce this picture.', space_after=8)
figure(f'{SHOTS}/06-savings-default.png', 6.5, 'Proof 1: CPS Savings Report, this month (console, today 2026-10-08 00:53 EET).')

# ---------- claim 2 ----------
heading('Claim 2 — On the first cluster, node autoscaling is configured and working — the low score has a different cause.')
p('Proof: the Savings Report of dema-platform-services. Realized savings +$2,930 this month, and under '
  '"Node autoscaler impact": baseline 684 CPU vs. 485 CPU actually provisioned — CAST AI removed ~199 CPU of waste. '
  'Node autoscaling has managed this cluster since July 22.', space_after=8)
figure(f'{SHOTS}/22-main-savings-last30d.png', 6.5, 'Proof 2a: dema-platform-services Savings Report (console, today).')
p('The 4.9 score (your screenshot showed 4.6 — it drifts) comes from remaining headroom, not from missing '
  'configuration: 0% Spot usage and a pending rebalancing recommendation. That headroom is also the next saving: '
  'Spot scenarios estimate −48.7% to −73.5% of the current ~$12.6k per month.', space_after=8)
figure(f'{SHOTS}/01-cluster-list.png', 6.5, 'Proof 2b: Cluster list, today — scores 4.9 / 6.2 side by side.')

# ---------- claim 3 ----------
heading('Claim 3 — The "couple of dozen Euro in minus" on the well-configured cluster is the Workload Autoscaler line — negative by definition, not by error.')
p('Documented formula (docs.cast.ai/docs/savings-baseline):', space_after=2)
p('workload autoscaler savings = (original requests − current requests) × max(current price, baseline price)',
  italic=True, size=9.5)
p('On both clusters the Workload Autoscaler raised requests that were set too low — a stability correction. '
  'Your test cluster recorded 348 OOM kills in the last 30 days; under-requested memory is a classic cause. '
  'When current requests are higher than the originals, this line is negative by construction. In August it was '
  '−$24.81 ≈ −€22 — exactly the figure you saw. Two things matter here:', space_after=4)
p('1. This value is already contained in Total savings. It is a sub-attribution, not an additional loss on top '
  '(documented behavior — Total savings = node autoscaler savings; the workload effect enters it via reduced demand).', space_after=2)
p('2. It turned positive on September 24, once rightsizing caught up.', space_after=8)
p('Proof: the same report for dema-platform-services-test, today — realized savings +$66 (45% under baseline), '
  'and the Workload Autoscaler line ≈ +$9, positive.', space_after=8)
figure(f'{SHOTS}/09-savings-main-only.png', 6.5, 'Proof 3: dema-platform-services-test Savings Report (console, today).')

# ---------- calculation ----------
heading('The calculation, compressed')
tbl = doc.add_table(rows=5, cols=3)
tbl.style = 'Table Grid'
rows = [
    ('Month', 'dema-platform-services — realized / workload line', 'dema-platform-services-test — realized / workload line'),
    ('Jul 22–31', '+$2,567 / −$17', '+$70 / −$6'),
    ('August', '+$9,951 / −$75', '+$213 / −$25 (≈ −€22 ← your figure)'),
    ('September', '+$11,690 / −$81', '+$202 / −$9'),
    ('October 1–8 (console today)', '+$2,930 / ≈ +$3', '+$66 / ≈ +$9'),
]
for i, (a, b, c) in enumerate(rows):
    for j, val in enumerate((a, b, c)):
        cell = tbl.rows[i].cells[j]
        cell.text = val
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(9)
            if i == 0: run.bold = True
doc.add_paragraph()
p('Verdict: nothing is broken on either cluster. The minus you see is the Workload Autoscaler sub-attribution for '
  'corrected under-sizing — stability money, not lost money. The real remaining lever is Spot adoption '
  '(−48.7% to −73.5% estimated on dema-platform-services). Say the word and we walk through it in 30 minutes.')

p('One note from our side, for completeness: the cloud-cost telemetry for both clusters has been stale since '
  'September 16 (some report tiles under-count as a result). We are correcting this — no action needed from you.')
p('Sources: Savings Report https://docs.cast.ai/docs/savings-report — methodology and formulas '
  'https://docs.cast.ai/docs/savings-baseline. Screenshots taken today, 2026-10-08, from your console; '
  'monthly history from the CAST AI reporting API (the console offers rolling time windows only).')
p('')
p('Best regards,')
p('Ebrahim Ramadan', bold=True, space_after=0)
p('CAST AI — ebrahim@cast.ai', size=9)

doc.save(OUT)
print('SAVED', OUT)
