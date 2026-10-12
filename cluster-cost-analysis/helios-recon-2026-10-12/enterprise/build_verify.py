#!/usr/bin/env python3
"""Verification layer: baseline params + rebalancing + inventory markers + economics -> classes, revised nets."""
import json, csv, glob, collections
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = '/Users/eramadan/castai/cluster-cost-analysis/helios-recon-2026-10-12'
ENT  = f'{BASE}/enterprise'
V    = f'{ENT}/verify'

bill = {e['entityName']: e['usage'] for e in json.load(open(f'{BASE}/billing_phase2_sep26_enterprise.json'))['detail']['entities']}
vr = {}
for fp in glob.glob(f'{ENT}/vr_*.json'):
    if '_slice' in fp or '_guard' in fp:
        continue
    for it in json.load(open(fp)).get('items', []):
        vr[it['clusterId']] = it
for it in json.load(open(f'{BASE}/vr_sep26_berlin_all_clusters.json'))['items']:
    vr[it['clusterId']] = it
cls = {}
for fp in glob.glob(f'{ENT}/classic_*.json'):
    u = fp.split('classic_')[1][:36]
    try:
        s = json.load(open(fp)).get('summary', {})
        cls[u] = float(s.get('totalSavings', 0) or 0)
    except Exception:
        pass
meta = {}
for fp in glob.glob(f'{ENT}/result_slice*.tsv'):
    for line in open(fp):
        c = line.rstrip('\n').split('\t')
        if len(c) < 6 or c[0].strip() == 'ORG_NAME':
            continue
        meta[c[3]] = c[0:6]

inv = {}
for fp in glob.glob(f'{ENT}/inv_*.json'):
    try:
        j = json.load(open(fp))
    except Exception:
        continue
    for it in j.get('items', []):
        inv[it['id']] = it

scope = {}
for line in open(f'{ENT}/classic_scope.tsv'):
    u, oid, org, name = line.rstrip('\n').split('\t')
    scope[u] = dict(oid=oid, org=org, name=name)

rows = []
for u, s in scope.items():
    m = meta.get(u)
    it = inv.get(u, {})
    bp = rp = None
    try:
        bp = json.load(open(f'{V}/bp_{u}.json'))
    except Exception:
        pass
    try:
        rp = json.load(open(f'{V}/rp_{u}.json'))
    except Exception:
        pass
    v = bill.get(u)
    vit = vr.get(u)
    phase2 = it.get('isPhase2')
    vsav = float(vit['cost']['totalSavings']) if vit else None
    fee5 = (v or 0) * 5.0 if v is not None else None
    fee57 = (v or 0) * 5.7 if v is not None else None
    net5 = (vsav - (fee5 or 0)) if vsav is not None else (None if not fee5 else -fee5)
    net57 = (vsav - (fee57 or 0)) if vsav is not None else (None if not fee57 else -fee57)
    # classification
    if phase2 is True:
        cl = 'A standard-managed'
    elif phase2 is False and (v or 0) > 0.01 and vsav is not None:
        cl = 'B enterprise/karpenter-style'
    elif phase2 is False and vsav is not None:
        cl = 'C unmanaged-with-VR'
    elif (v or 0) > 0.01:
        cl = 'D billed-no-data'
    else:
        cl = 'E ghost/disconnected'
    rpdates = []
    if rp:
        rpdates = [i.get('createdAt', '') for i in rp.get('items', []) if i.get('createdAt')]
    rows.append(dict(u=u, org=s['org'], name=s['name'], region=(m[5] if m else '-'),
                     phase2=phase2, managedBy=it.get('managedBy') or '', status=it.get('status', '?'),
                     createdAt=(it.get('createdAt') or '')[:10], firstOp=(it.get('firstOperationAt') or '')[:19].replace('T', ' '),
                     btype=(bp or {}).get('baselineType', ''), bstart=(bp or {}).get('baselinePeriodStartTime', '')[:10],
                     bend=(bp or {}).get('baselinePeriodEndTime', '')[:10], brecomp=(bp or {}).get('updateTime', '')[:10],
                     cpuf=round(float((bp or {}).get('cpuOverprovisioningFactor', 0) or 0), 4),
                     ramf=round(float((bp or {}).get('memoryOverprovisioningFactor', 0) or 0), 4),
                     firstRebalance=(min(rpdates) if rpdates else '')[:10], nPlans=len(rpdates),
                     vcpu=v, vsav=vsav, csav=cls.get(u), fee5=fee5, fee57=fee57, net5=net5, net57=net57, cl=cl))

# csv
with open(f'{V}/verification_full.csv', 'w', newline='') as fh:
    w = csv.writer(fh)
    w.writerow(['class', 'org', 'cluster_name', 'uuid', 'region', 'isPhase2', 'status', 'created', 'first_operation',
                'first_rebalance', 'rebalance_plans', 'baseline_type', 'baseline_start', 'baseline_end', 'baseline_recomputed',
                'cpu_factor', 'ram_factor', 'billable_vcpu', 'vr_savings_usd', 'classic_savings_usd',
                'net_vr_usd_at_5.00', 'net_vr_usd_at_5.70'])
    for r in sorted(rows, key=lambda r: (r['cl'], r['org'].lower(), -(r['vsav'] or 0))):
        w.writerow([r['cl'], r['org'], r['name'], r['u'], r['region'], r['phase2'], r['status'], r['createdAt'],
                    r['firstOp'], r['firstRebalance'], r['nPlans'], r['btype'], r['bstart'], r['bend'], r['brecomp'],
                    r['cpuf'] or '', r['ramf'] or '',
                    '' if r['vcpu'] is None else f"{r['vcpu']:.3f}",
                    '' if r['vsav'] is None else f"{r['vsav']:.2f}",
                    '' if r['csav'] is None else f"{r['csav']:.2f}",
                    '' if r['net5'] is None else f"{r['net5']:.2f}",
                    '' if r['net57'] is None else f"{r['net57']:.2f}"])

# summary by class
bycl = collections.OrderedDict()
for r in rows:
    d = bycl.setdefault(r['cl'], dict(n=0, v=0, vsav=0, fee5=0, net5=0, net57=0, names=[]))
    d['n'] += 1
    d['v'] += r['vcpu'] or 0
    d['vsav'] += r['vsav'] or 0
    d['fee5'] += r['fee5'] or 0
    d['net5'] += r['net5'] or 0
    d['net57'] += r['net57'] or 0
    d['names'].append(f"{r['org']}/{r['name']}")
fmt = lambda x, d=0: f"{x:,.{d}f}"
print(f"{'class':<28}{'n':>4}{'vCPU':>10}{'VR savings':>12}{'net@5':>12}{'net@5.7':>12}")
tot = dict(n=0, v=0, vsav=0, net5=0, net57=0)
for cl, d in bycl.items():
    print(f"{cl:<28}{d['n']:>4}{fmt(d['v'],1):>10}{fmt(d['vsav']):>12}{fmt(d['net5']):>12}{fmt(d['net57']):>12}")
    for k in tot:
        if k == 'n':
            tot['n'] += d['n']
        else:
            tot[k] += d[k]
print(f"{'Σ ALL':<28}{tot['n']:>4}{fmt(tot['v'],1):>10}{fmt(tot['vsav']):>12}{fmt(tot['net5']):>12}{fmt(tot['net57']):>12}")

# scenarios
a_pos5 = sum(r['net5'] for r in rows if r['cl'].startswith('A') and (r['net5'] or 0) > 0)
a_pos57 = sum(r['net57'] for r in rows if r['cl'].startswith('A') and (r['net57'] or 0) > 0)
a_all5 = sum(r['net5'] or 0 for r in rows if r['cl'].startswith('A'))
a_all57 = sum(r['net57'] or 0 for r in rows if r['cl'].startswith('A'))
sim = [r for r in rows if r['name'] == 'ngm-sim-eks'][0]
gsw = [r for r in rows if r['org'] == 'SI GSW CLO']
gsw5 = sum(r['net5'] or 0 for r in gsw)
gsw_clean5 = gsw5 - (sim['net5'] or 0) - sum(r['net5'] or 0 for r in gsw if r['cl'].startswith('C'))
print()
print(f"positive-net-only Σ @5.00  ALL classes : {fmt(sum(r['net5'] for r in rows if (r['net5'] or 0)>0))} (clusters: {sum(1 for r in rows if (r['net5'] or 0)>0)})")
print(f"A-class Σ all rows  @5.00: {fmt(a_all5)} | @5.70: {fmt(a_all57)}")
print(f"A-class positives   @5.00: {fmt(a_pos5)} | @5.70: {fmt(a_pos57)}")
print(f"A-class minus ngm-sim-eks @5.00: {fmt(a_all5 - (sim['net5'] or 0))}")
print(f"SI GSW CLO Σ net@5: {fmt(gsw5)} | without sim(+C classes): {fmt(gsw_clean5)}")

# xlsx sheet
wb = load_workbook(f'{BASE}/ENTERPRISE-SEP26-NET-SAVINGS.xlsx')
if 'Baseline & onboarding' in wb.sheetnames:
    del wb['Baseline & onboarding']
ws = wb.create_sheet('Baseline & onboarding')
ws.merge_cells('A1:S1')
ws['A1'] = 'Cluster classification & baseline/onboarding evidence (for CAST AI engineering call 2026-10-13) — A=standard-managed · B=enterprise/karpenter-style · C=unmanaged-with-VR · D=billed-no-data · E=ghost/disconnected'
ws['A1'].font = Font(bold=True, size=11)
ws.append([])
hdr = ['Class', 'Organization', 'Cluster name', 'CAST AI cluster ID', 'Region', 'isPhase2', 'Status',
       'Created', 'First CAST op', 'First rebalance', 'Plans', 'Baseline type', 'Baseline start', 'Baseline end',
       'Recomputed', 'CPU factor', 'RAM factor', 'VR savings $', 'Net VR @5.00']
ws.append(hdr)
for r in sorted(rows, key=lambda r: (r['cl'], r['org'].lower(), -(r['vsav'] or 0))):
    ws.append([r['cl'], r['org'], r['name'], r['u'], r['region'],
               'yes' if r['phase2'] else ('no' if r['phase2'] is False else '?'), r['status'],
               r['createdAt'], r['firstOp'], r['firstRebalance'], r['nPlans'] or None,
               r['btype'], r['bstart'], r['bend'], r['brecomp'],
               r['cpuf'] or None, r['ramf'] or None, r['vsav'], r['net5']])
for cell in ws[3]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F3864')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
clfill = {'A': 'C6EFCE', 'B': 'FFF2CC', 'C': 'FCE4D6', 'D': 'D9D9D9', 'E': 'D9D9D9'}
for row in ws.iter_rows(min_row=4):
    k = str(row[0].value or '')[:1]
    if k in clfill:
        for c in row:
            c.fill = PatternFill('solid', fgColor=clfill[k])
    for c in row:
        c.border = Border(bottom=Side(style='thin', color='BFBFBF'))
        if c.column in (16, 17) and isinstance(c.value, float):
            c.number_format = '0.0000'
        if c.column in (18, 19) and isinstance(c.value, (int, float)):
            c.number_format = '#,##0'
            if c.column == 19 and c.value < 0:
                c.font = Font(color='C00000')
ws.freeze_panes = 'A4'
ws.auto_filter.ref = f'A3:S{3 + len(rows)}'
widths = [26, 30, 34, 38, 12, 8, 8, 11, 18, 12, 7, 17, 11, 11, 11, 10, 10, 13, 13]
for i, wdt in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = wdt
wb.save(f'{BASE}/ENTERPRISE-SEP26-NET-SAVINGS.xlsx')
print('\nxlsx sheet added')
