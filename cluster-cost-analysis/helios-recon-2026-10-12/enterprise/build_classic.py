#!/usr/bin/env python3
"""Classic-savings merge: 49 classic JSONs + billing fees -> md table + xlsx sheet."""
import json, csv, glob, collections
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = '/Users/eramadan/castai/cluster-cost-analysis/helios-recon-2026-10-12'
ENT  = f'{BASE}/enterprise'

def fmt(x, d=0):
    return f"{x:,.{d}f}" if x is not None else "—"

bill = {e['entityName']: e['usage'] for e in json.load(open(f'{BASE}/billing_phase2_sep26_enterprise.json'))['detail']['entities']}
meta = {}
for fp in glob.glob(f'{ENT}/result_slice*.tsv'):
    for line in open(fp):
        c = line.rstrip('\n').split('\t')
        if len(c) < 6 or c[0].strip() == 'ORG_NAME':
            continue
        meta[c[3].strip()] = (c[0].strip(), c[1].strip(), c[2].strip(), c[4].strip(), c[5].strip())
oid2name = {l.split('\t')[0]: l.split('\t')[1].strip() for l in open(f'{BASE}/orgs_130.tsv')}

scope = {}
for line in open(f'{ENT}/classic_scope.tsv'):
    u, oid, org, name = line.rstrip('\n').split('\t')
    scope[u] = (oid, org, name)

rows, errs = [], []
for fp in glob.glob(f'{ENT}/classic_*.json'):
    u = fp.split('classic_')[1][:36]
    try:
        j = json.load(open(fp))
    except Exception:
        continue
    s = j.get('summary', {})
    items = j.get('items', []) or []
    oid, org, name = scope.get(u, ('?', '?', u))
    down = sum(float(it.get('downscalingSavings', 0) or 0) for it in items)
    spot = sum(float(it.get('spotSavings', 0) or 0) for it in items)
    rows.append(dict(org=org, oid=oid, name=name, uuid=u,
                     cost=float(s.get('totalCost', 0) or 0),
                     sav=float(s.get('totalSavings', 0) or 0),
                     down=down, spot=spot))

# error set from scope (cluster files missing => ERR)
have = {r['uuid'] for r in rows}
for u in sorted(set(scope) - have):
    oid, org, name = scope[u]
    v = bill.get(u)
    errs.append(dict(org=org, name=name, uuid=u, vcpu=v, fee5=(v or 0)*5 if v else None,
                     reason='classic API 400: cluster is read-only (not CAST-managed at cluster level)'))

# join fees
for r in rows:
    v = bill.get(r['uuid'])
    m = meta.get(r['uuid'])
    r['acct'] = m[3] if m else '-'
    r['region'] = m[4] if m else '-'
    r['vcpu'] = v
    r['fee5'], r['fee57'] = (v or 0)*5.00, (v or 0)*5.70
    r['net5'], r['net57'] = r['sav'] - r['fee5'], r['sav'] - r['fee57']

orgs = collections.OrderedDict()
for r in sorted(rows, key=lambda r: (r['org'].lower(), -r['sav'])):
    orgs.setdefault(r['org'], []).append(r)

# ---------- md ----------
HDR = "| Cluster name | Cluster ID (name#account#region) | CAST AI cluster ID | Region | Classic cost $ | Downscaling $ | Spot $ | **Classic gross savings $** | Billable vCPU | Fee $ @5.00 | Fee $ @5.70 | **Net $ @5.00** | **Net $ @5.70** |"
DIV = "|---|---|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|"
L = ["# Siemens AG enterprise — CLASSIC savings API, September 2026, fees & net",
     "",
     "**Basis:** `GET /v1/cost-reports/clusters/{id}/savings?startTime=2026-09-01T00:00:00Z&endTime=2026-10-01T00:00:00Z&stepSeconds=86400` per cluster · fees = enterprise billing (phase2) × $5.00/$5.70. 71 clusters in scope (56 billed ∪ VR rows); **49 returned classic data; 22 answered 400 'cluster is read-only'** (not CAST-managed at cluster level — listed at bottom; their fees are NOT netted into the classic totals since classic gives them no savings number).",
     ""]
T = collections.defaultdict(float)
roll = []
for org, rs in orgs.items():
    st = collections.defaultdict(float)
    for r in rs:
        for k in ('cost', 'sav', 'down', 'spot', 'fee5', 'fee57', 'net5', 'net57'):
            st[k] += r[k]
        st['v'] += r['vcpu'] or 0
    roll.append((org, len(rs), st))
    for k, val in st.items():
        T[k] += val
    T['clu'] += len(rs)
L += ["## Organization roll-up (classic basis)",
      "",
      "| Organization | Clusters w/ classic data | Classic cost $ | Classic gross savings $ | Billable vCPU | Fee $ @5.00 | Fee $ @5.70 | **Net $ @5.00** | **Net $ @5.70** |",
      "|---|---:|---:|---:|---:|---:|---:|---:|---:|"]
for org, n, st in sorted(roll, key=lambda x: -x[2]['sav']):
    L.append(f"| {org} | {n} | {fmt(st['cost'])} | {fmt(st['sav'])} | {fmt(st['v'],2)} | {fmt(st['fee5'],2)} | {fmt(st['fee57'],2)} | **{'+' if st['net5']>=0 else ''}{fmt(st['net5'])}** | **{'+' if st['net57']>=0 else ''}{fmt(st['net57'])}** |")
L.append(f"| **Σ ENTERPRISE (classic)** | **{T['clu']}** | **{fmt(T['cost'])}** | **{fmt(T['sav'])}** | **{fmt(T['v'],2)}** | **{fmt(T['fee5'],2)}** | **{fmt(T['fee57'],2)}** | **{'+' if T['net5']>=0 else ''}{fmt(T['net5'])}** | **{'+' if T['net57']>=0 else ''}{fmt(T['net57'])}** |")
L.append("")
for org, rs in orgs.items():
    L.append(f"## {org}")
    L.append("")
    L.append(HDR); L.append(DIV)
    for r in rs:
        cid = f"{r['name']}#{r['acct']}#{r['region']}" if r['acct'] != '-' else r['uuid']
        L.append(f"| {r['name']} | {cid} | `{r['uuid']}` | {r['region']} | {fmt(r['cost'])} | {fmt(r['down'])} | {fmt(r['spot'])} | **{fmt(r['sav'])}** | {fmt(r['vcpu'],2) if r['vcpu'] is not None else '—'} | {fmt(r['fee5'],2)} | {fmt(r['fee57'],2)} | {fmt(r['net5'])} | {fmt(r['net57'])} |")
    L.append("")
L += ["## Clusters with no classic data (400 read-only)",
      "",
      "| Organization | Cluster | CAST AI cluster ID | Billable vCPU | Fee $ @5.00 | Reason |",
      "|---|---|---|---:|---:|---|"]
for e in sorted(errs, key=lambda e: (e['org'], e['name'])):
    L.append(f"| {e['org']} | {e['name']} | `{e['uuid']}` | {fmt(e['vcpu'],2) if e['vcpu'] is not None else '—'} | {fmt(e['fee5'],2) if e['fee5'] is not None else '—'} | {e['reason']} |")
err_fee = sum(e['fee5'] or 0 for e in errs)
err_v = sum(e['vcpu'] or 0 for e in errs)
L += ["", f"*These {len(errs)} clusters carry {fmt(err_v,2)} billable vCPU (fee Σ ${fmt(err_fee,2)} @5.00) with no classic savings figure — do not compare the classic Σ net against the VR Σ net without accounting for this exclusion (SMO Railigent X alone carries ~1,073 vCPU / $5.4k fee here while holding $158.7k of VR savings).*",
      "",
      "## Notes",
      "",
      "1. Classic ≠ VR basis: classic is capacity/mechanism metered (downscaling + spot), can back-fill savings before a cluster's real CAST onboarding, and has no demand cap. Use VR for leadership KPI; classic for mechanism/cash narrative.",
      "2. Cluster classic_N/A rows: cluster is in a read-only onboarding state for the classic endpoint (e.g. externally-managed or woop-only). VR is the only savings basis available for them.",
      f"3. Savings split: downscaling ${fmt(T['down'])} ({T['down']/T['sav']*100:.0f}%) + spot ${fmt(T['spot'])} ({T['spot']/T['sav']*100:.0f}%) — downscaling-dominated, but spot is real at enterprise level (SI EA electrificationx, HAFAS H&O).",
      "4. Fees applied USD per instruction; if contract is EUR, ×~1.15.",
      "",
      "Provenance: enterprise/classic_<clusterID>.json (49 pulls, parallel xargs, 2026-10-12) + enterprise/result_slice{1..8}.tsv + billing_phase2_sep26_enterprise.json.",
      ""]
open(f'{BASE}/CLASSIC-SEP26-NET-SAVINGS.md', 'w').write("\n".join(L))

# ---------- csv ----------
with open(f'{BASE}/classic_sep26_enterprise_net_savings.csv', 'w', newline='') as fh:
    w = csv.writer(fh)
    w.writerow(['organization', 'cluster_name', 'cluster_id', 'castai_cluster_id', 'region',
                'classic_cost_usd', 'downscaling_usd', 'spot_usd', 'classic_savings_usd',
                'billable_vcpu_phase2', 'fee_usd_at_5.00', 'fee_usd_at_5.70',
                'net_usd_at_5.00', 'net_usd_at_5.70'])
    for org, rs in orgs.items():
        for r in rs:
            cid = f"{r['name']}#{r['acct']}#{r['region']}" if r['acct'] != '-' else r['uuid']
            w.writerow([org, r['name'], cid, r['uuid'], r['region'], f"{r['cost']:.2f}", f"{r['down']:.2f}",
                        f"{r['spot']:.2f}", f"{r['sav']:.2f}",
                        '' if r['vcpu'] is None else f"{r['vcpu']:.3f}",
                        f"{r['fee5']:.2f}", f"{r['fee57']:.2f}", f"{r['net5']:.2f}", f"{r['net57']:.2f}"])

# ---------- xlsx sheet ----------
HDR_FILL = PatternFill('solid', fgColor='1F3864')
TOT_FILL = PatternFill('solid', fgColor='FFE699')
THIN = Border(bottom=Side(style='thin', color='BFBFBF'))
wb = load_workbook(f'{BASE}/ENTERPRISE-SEP26-NET-SAVINGS.xlsx')
if 'Classic savings' in wb.sheetnames:
    del wb['Classic savings']
ws = wb.create_sheet('Classic savings')
ws.merge_cells('A1:M1')
ws['A1'] = 'Classic savings API basis — September 2026 (49 clusters with classic data; 22 more are classic-N/A, see md). Fee rates $5.00/$5.70 per billable vCPU.'
ws['A1'].font = Font(bold=True, size=12)
ws.append([])
hdr = ['Organization', 'Cluster name', 'Cluster ID (name#account#region)', 'CAST AI cluster ID', 'Region',
       'Classic cost $', 'Downscaling $', 'Spot $', 'Classic gross savings $', 'Billable vCPU',
       'Fee $ @5.00', 'Fee $ @5.70', 'Net $ @5.00']
ws2hdr = hdr + ['Net $ @5.70']
ws.append(ws2hdr)
for org, rs in orgs.items():
    for r in rs:
        cid = f"{r['name']}#{r['acct']}#{r['region']}" if r['acct'] != '-' else r['uuid']
        ws.append([org, r['name'], cid, r['uuid'], r['region'], r['cost'], r['down'], r['spot'], r['sav'],
                   r['vcpu'], r['fee5'], r['fee57'], r['net5'], r['net57']])
for cell in ws[3]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=4):
    for c in row:
        c.border = THIN
        if c.column in (6, 7, 8, 9) and isinstance(c.value, (int, float)):
            c.number_format = '#,##0'
        if c.column in (11, 12, 13, 14) and isinstance(c.value, (int, float)):
            c.number_format = '#,##0.00'
            if c.column >= 13 and c.value < 0:
                c.font = Font(color='C00000')
        if c.column == 10 and isinstance(c.value, (int, float)):
            c.number_format = '#,##0.000'
ws.freeze_panes = 'A4'
ws.auto_filter.ref = f'A3:N{3 + sum(len(rs) for rs in orgs.values())}'
names = 'BCDEFGHIJKLMN'
for i, width in enumerate([30, 30, 46, 38, 12, 13, 13, 10, 15, 12, 12, 12, 13, 13], start=1):
    ws.column_dimensions[get_column_letter(i)].width = width
wb.save(f'{BASE}/ENTERPRISE-SEP26-NET-SAVINGS.xlsx')

print('md + csv + xlsx sheet written')
print(f"classic Σ: cost {fmt(T['cost'])}  gross sav {fmt(T['sav'])}  down {fmt(T['down'])}  spot {fmt(T['spot'])}  vCPU {fmt(T['v'],2)}")
print(f"fee @5 {fmt(T['fee5'],2)} | net @5 {fmt(T['net5'])}  net @5.7 {fmt(T['net57'])}")
print(f"classic-N/A clusters: {len(errs)} = {fmt(err_v,2)} vCPU, fee @5 ${fmt(err_fee,2)}")
for org, n, st in sorted(roll, key=lambda x: -x[2]['sav']):
    print(f"  {org:<34} {n:>3} clusters  sav {fmt(st['sav']):>10}  net@5 {fmt(st['net5']):>10}")
