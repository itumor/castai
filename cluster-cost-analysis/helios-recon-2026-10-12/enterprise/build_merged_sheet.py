#!/usr/bin/env python3
"""Merge baseline + VR(all clusters) + classic into one sheet in the user's workbook copy."""
import csv, glob
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = '/Users/eramadan/castai/cluster-cost-analysis/helios-recon-2026-10-12'
TARGETS = ['/Users/eramadan/Downloads/ENTERPRISE-SEP26-NET-SAVINGS (1).xlsx',
           f'{BASE}/ENTERPRISE-SEP26-NET-SAVINGS.xlsx']
SHEET = 'Merged — all data'

def num(v):
    try:
        f = float(v)
        return f
    except (TypeError, ValueError):
        return None

vr = {}
for r in csv.DictReader(open(f'{BASE}/vr_sep26_enterprise_net_savings.csv')):
    vr[r['castai_cluster_id']] = r
verify = {}
for r in csv.DictReader(open(f'{BASE}/enterprise/verify/verification_full.csv')):
    verify[r['uuid']] = r
cls = {}
for r in csv.DictReader(open(f'{BASE}/classic_sep26_enterprise_net_savings.csv')):
    cls[r['castai_cluster_id']] = r

CL_ORDER = {'A': 0, 'B': 1, 'C': 2, 'D': 3, 'E': 4}
keys = set(vr) | set(verify) | set(cls)

def sort_key(u):
    v = verify.get(u, {})
    c0 = (v.get('class') or 'E')[0]
    return (CL_ORDER.get(c0, 5), (v.get('org') or vr.get(u, {}).get('organization') or '').lower(),
            -(num(vr.get(u, {}).get('gross_savings_usd')) or 0))

rows = sorted(keys, key=sort_key)
HDR = ['Class', 'Organization', 'Cluster name', 'CAST AI cluster ID', 'Region', 'isPhase2', 'Status', 'Created',
       'First CAST op', 'First rebalance', 'Plans', 'Baseline type', 'Baseline start', 'Baseline end', 'Recomputed',
       'CPU factor', 'RAM factor', 'Cluster ID (name#account#region)',
       'Actual $ (VR)', 'Projected $ (VR)', 'VR savings $', 'Net VR @5.00', 'Net VR @5.70',
       'Classic cost $', 'Downscaling $', 'Spot $', 'Classic gross savings $', 'Net classic @5.00', 'Net classic @5.70',
       'Billable vCPU', 'Fee $ @5.00', 'Fee $ @5.70', 'Has VR row']

def build_row(u):
    v, c, k = verify.get(u, {}), vr.get(u, {}), cls.get(u, {})
    cls_kv = {x: num(k[x]) for x in ('classic_cost_usd', 'downscaling_usd', 'spot_usd', 'classic_savings_usd',
                                     'net_usd_at_5.00', 'net_usd_at_5.70')} if k else {}
    return [
        (v.get('class') or '').split(' ')[0].upper() if v else '',
        v.get('org') or c['organization'],
        v.get('cluster_name') or c['cluster_name'],
        u,
        v.get('region') or c.get('region', ''),
        v.get('isPhase2') if v.get('isPhase2') not in ('', None) else '',
        v.get('status', ''), v.get('created', ''), v.get('first_operation', ''),
        v.get('first_rebalance', ''), num(v.get('rebalance_plans')),
        v.get('baseline_type', ''), v.get('baseline_start', ''), v.get('baseline_end', ''), v.get('baseline_recomputed', ''),
        num(v.get('cpu_factor')), num(v.get('ram_factor')),
        c.get('cluster_id', u),
        num(c.get('actual_usd')), num(c.get('projected_usd')), num(c.get('gross_savings_usd')),
        num(c.get('net_usd_at_5.00')), num(c.get('net_usd_at_5.70')),
        cls_kv.get('classic_cost_usd'), cls_kv.get('downscaling_usd'), cls_kv.get('spot_usd'),
        cls_kv.get('classic_savings_usd'), cls_kv.get('net_usd_at_5.00'), cls_kv.get('net_usd_at_5.70'),
        num(c.get('billable_vcpu_phase2')), num(c.get('fee_usd_at_5.00')), num(c.get('fee_usd_at_5.70')),
        c.get('has_vr', ''),
    ]

for T in TARGETS:
    try:
        wb = load_workbook(T)
    except Exception as e:
        print(f'skip {T}: {e}')
        continue
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR))
    ws.cell(row=1, column=1).value = ('One merged view — baseline & onboarding + Value-Realization economics + Classic savings + fees '
                                      '(join key: CAST AI cluster ID). Classes: A standard-managed · B enterprise/karpenter-style · '
                                      'C unmanaged-with-VR · D billed-no-data · E ghost/disconnected. September 2026.')
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 30
    ws.append(HDR)
    for u in rows:
        ws.append(build_row(u))
    bands = {'1': '595959', '12': '70AD47', '19': '2E75B6', '24': '833C00', '30': '1F4E79'}
    for cell in ws[2]:
        col = cell.column
        fill = '595959' if col <= 11 else ('70AD47' if col <= 18 else ('2E75B6' if col <= 23 else ('833C00' if col <= 29 else '1F4E79')))
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 32
    clfill = {'A': 'C6EFCE', 'B': 'FFF2CC', 'C': 'FCE4D6', 'D': 'D9D9D9', 'E': 'D9D9D9'}
    for row in ws.iter_rows(min_row=3):
        k = str(row[0].value or '')[:1]
        for c in row:
            c.border = Border(bottom=Side(style='thin', color='BFBFBF'))
            if c.column == 1 and k in clfill:
                c.fill = PatternFill('solid', fgColor=clfill[k])
            if c.column in (16, 17) and isinstance(c.value, float):
                c.number_format = '0.0000'
            if c.column in (19, 20, 21, 24, 25, 26, 27) and isinstance(c.value, (int, float)):
                c.number_format = '#,##0'
            if c.column in (22, 23, 28, 29, 31, 32) and isinstance(c.value, (int, float)):
                c.number_format = '#,##0.00'
                if c.value < 0:
                    c.font = Font(color='C00000')
            if c.column == 30 and isinstance(c.value, (int, float)):
                c.number_format = '#,##0.000'
    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(HDR))}{2 + len(rows)}'
    widths = [7, 22, 30, 36, 12, 8, 8, 11, 18, 12, 6, 15, 11, 11, 11, 10, 10, 44,
              13, 13, 13, 12, 12, 13, 12, 9, 15, 13, 13, 12, 12, 12, 9]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(T)
    print(f'updated {T}: {len(rows)} rows x {len(HDR)} cols')
